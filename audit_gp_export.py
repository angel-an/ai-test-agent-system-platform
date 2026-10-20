# -*- coding: utf-8 -*-
"""导出商品 / 支付两模块的测试完成情况表格（可直接回填/交接）。"""
from __future__ import annotations

import json
from pathlib import Path

import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path("out_yl_case_audit")
SUMMARY = OUT / "gp_summary.json"
TARGET = Path("商品支付两模块测试完成情况-20261017.xlsx")
PROJ = "942226d9-6071-43fa-b2bb-f7cbc93a0028"
MODULES = ("商品主數據", "YL-主數據管理-商品", "支付方式主數據", "YL-主數據管理-支付方式")
LABEL = {
    "商品主數據": "商品·文档生成",
    "YL-主數據管理-商品": "商品·人工导入",
    "支付方式主數據": "支付·文档生成",
    "YL-主數據管理-支付方式": "支付·人工导入",
}
VERDICT_CN = {
    "PASSED": "通过",
    "PARTIAL": "部分通过",
    "UNVERIFIED": "UI层无法判定",
    "FAILED_ASSERT": "断言失败",
    "BLOCKED_EXEC": "执行阻断",
    "FAILED_EXEC": "执行失败",
    "NOT_EXECUTED": "无可执行断言",
    "SKIPPED_NO_CHECK": "编译期无检查",
    "NOT_RUN": "未执行",
    "TIMEOUT": "超时",
    "ERROR": "调用异常",
}

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_sheet(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> None:
    data = json.loads(SUMMARY.read_text(encoding="utf-8"))
    cases = data["cases"]
    totals = data["totals"]
    layer = data["assertion_layer"]

    # 补全无脚本用例（不在资产范围内，但属于模块测试范围）
    conn = psycopg2.connect(host="127.0.0.1", port=5432, user="postgres",
                            password="postgres", dbname="langgraph_platform_db")
    cur = conn.cursor()
    cur.execute("""
    select t.module, t.identifier, t.name, coalesce(t.priority::text,''),
           coalesce(t.custom_fields->'duplicate_of'->>0, t.custom_fields->>'duplicate_of',''),
           length(coalesce(t.description,''))
    from test_cases t
    where t.project_id = %s and t.module = any(%s)
      and not exists (select 1 from web_tests w where w.test_case_id = t.id
                      and w.generated_by_agent = 'web_case_link_service')
    order by t.module, t.identifier
    """, (PROJ, list(MODULES)))
    no_asset = cur.fetchall()
    # 用例名映射（含无脚本）
    cur.execute("select identifier, name from test_cases where project_id = %s", (PROJ,))
    names = dict(cur.fetchall())
    conn.close()

    wb = Workbook()

    # ---- Sheet1 汇总 ----
    ws = wb.active
    ws.title = "汇总"
    ws.append(["模块", "脚本资产", "通过", "部分通过", "UI层无法判定", "断言失败",
               "执行阻断", "无可执行断言", "编译期无检查"])
    from collections import Counter
    by_mod = {}
    for c in cases:
        by_mod.setdefault(c["module"], Counter())[c["verdict"]] += 1
    for mod in MODULES:
        cnt = by_mod.get(mod)
        if not cnt:
            continue
        ws.append([LABEL[mod], sum(cnt.values()), cnt.get("PASSED",0), cnt.get("PARTIAL",0), cnt.get("UNVERIFIED",0),
                   cnt.get("FAILED_ASSERT", 0), cnt.get("BLOCKED_EXEC", 0),
                   cnt.get("NOT_EXECUTED", 0), cnt.get("SKIPPED_NO_CHECK", 0)])
    ws.append(["合计", sum(totals.values()), totals.get("PASSED", 0), totals.get("PARTIAL", 0),
               totals.get("UNVERIFIED", 0), totals.get("FAILED_ASSERT", 0),
               totals.get("BLOCKED_EXEC", 0), totals.get("NOT_EXECUTED", 0),
               totals.get("SKIPPED_NO_CHECK", 0)])
    ws.append([])
    ws.append(["断言层", f"检查项 {layer['plan_checks']}", f"通过 {layer['pass']}",
               f"无法判定 {layer['unverified']}", f"阻断 {layer['blocked']}",
               f"归因断言 {layer['attributed']}", f"配置级检查 {layer['config_only']}"])
    ws.append([])
    ws.append(["模块测试范围", "全部用例", "有脚本", "无脚本（需人工/补模板）", "其中重复标记"])
    total_cases = 0
    for mod in MODULES:
        mod_cases = [c for c in cases if c["module"] == mod]
        dup = sum(1 for r in no_asset if r[0] == mod and r[4])
        na = sum(1 for r in no_asset if r[0] == mod)
        total_cases += len(mod_cases) + na
        ws.append([LABEL[mod], len(mod_cases) + na, len(mod_cases), na, dup])
    ws.append(["合计", total_cases, len(cases), len(no_asset),
               sum(1 for r in no_asset if r[4])])
    style_sheet(ws, [16, 12, 10, 10, 14, 10, 10, 14, 14])

    # ---- Sheet2 用例明细 ----
    ws2 = wb.create_sheet("用例明细")
    ws2.append(["模块", "用例编号", "用例名称", "优先级", "资产编号", "编译状态",
                "可执行检查", "执行结论", "通过/检查", "归因断言", "配置级检查", "执行记录", "说明"])
    for c in sorted(cases, key=lambda x: (MODULES.index(x["module"]), x["case"])):
        ws2.append([
            LABEL[c["module"]], c["case"], names.get(c["case"], ""), c["priority"] or "",
            c["asset"] or "", c["compile_status"],
            c["form_executable_checks"] if c["form_executable_checks"] is not None else "",
            VERDICT_CN.get(c["verdict"], c["verdict"]),
            f"{c['pass']}/{c['plan_checks']}" if c["pass"] is not None else "",
            c["attributed_checks"] if c["attributed_checks"] is not None else "",
            c["config_only_checks"] if c["config_only_checks"] is not None else "",
            c["run"] or "", (c["reason"] or "")[:200],
        ])
    for mod, ident, name, prio, dup, dlen in no_asset:
        ws2.append([LABEL[mod], ident, name, prio, "", "无脚本资产", "", "无脚本（需人工执行）",
                    "", "", "", "", f"重复用例（{dup}）" if dup else
                    ("用例含可人工执行步骤" if dlen > 50 else "用例无步骤文本，需补规则")])
    style_sheet(ws2, [16, 14, 42, 9, 14, 14, 10, 16, 11, 10, 10, 12, 46])

    # ---- Sheet3 失败与阻断 ----
    ws3 = wb.create_sheet("失败与阻断")
    ws3.append(["模块", "用例", "结论", "通过/检查", "原因", "断言观察"])
    for c in cases:
        if c["verdict"] not in ("FAILED_ASSERT", "BLOCKED_EXEC", "FAILED_EXEC", "PARTIAL", "UNVERIFIED", "TIMEOUT", "ERROR"):
            continue
        obs = " | ".join(f"{a.get('name')}: {a.get('observed')}"
                         for a in (c["assertions"] or [])
                         if a.get("outcome") in ("fail", "blocked", "unverified"))[:600]
        ws3.append([LABEL[c["module"]], c["case"], VERDICT_CN.get(c["verdict"], c["verdict"]),
                    f"{c['pass']}/{c['plan_checks']}" if c["pass"] is not None else "",
                    (c["reason"] or "")[:200], obs])
    style_sheet(ws3, [16, 14, 14, 11, 40, 70])

    # ---- Sheet4 需人工执行 ----
    ws4 = wb.create_sheet("需人工执行")
    ws4.append(["模块", "用例", "优先级", "缺口类型", "说明"])
    for c in cases:
        if c["verdict"] == "SKIPPED_NO_CHECK":
            ws4.append([LABEL[c["module"]], c["case"], c["priority"] or "",
                        "编译期无检查", f"pending={c['pending_locator_count']} skipped={c['skipped_assertion_count']}"])
    for mod, ident, name, prio, dup, dlen in no_asset:
        if dup:
            continue
        ws4.append([LABEL[mod], ident, prio, "无脚本资产",
                    "用例含可执行步骤，可人工执行" if dlen > 50 else "缺步骤文本与规则，需补用例"])
    style_sheet(ws4, [16, 14, 9, 18, 50])

    wb.save(TARGET)
    print(f"[xlsx] {TARGET}  用例明细 {ws2.max_row-1} 行")


if __name__ == "__main__":
    main()
