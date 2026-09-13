"""美心 Excel API 接口文档解析器。

读取固定模板（如 "YX CRM API Spec For Non POS"）的 .xlsx 文件，提取：
- 导读 Sheet 的公共规则（model/method/appid/ts/sign + MD5 签名算法）
- 每个接口 Sheet 的入参/出参字段、必填标记、类型、长度、示例、差异说明、疑问清单
- 每个单元格的来源位置（sheet / 行号 / 列号），供溯源与待确认项定位

设计原则：
- 只读取与结构化，不做业务推断；冲突/缺失交给 api_spec_normalizer 识别。
- 保留合并单元格展开后的坐标，保证"来源位置"可追溯。
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)


# 表头列名 -> 标准字段（同义词归一）
_HEADER_ALIASES = {
    "field": "field",
    "field层级": "field_level",
    "field層級": "field_level",
    "field remark": "field_remark",
    "匹配": "match",
    "說明/差異": "diff",
    "说明/差异": "diff",
    "follow up": "follow_up",
    "疑問清單": "questions",
    "疑问清单": "questions",
}

# 入参/出参/返回码分区标记（A 列）
#
# 注意：比对前会对单元格文本做 ``lower().replace(" ", "")`` 归一化，
# 因此这里的常量必须是**去空格小写**形式，否则 startswith 恒为 False
# （历史缺陷：_RETURN_CODE_MARK 曾写作 "return code"，导致返回码区从未被解析）。
_SECTION_INPUT = "field（入參）"
_SECTION_OUTPUT = "field（出參）"
_RETURN_CODE_MARK = "returncode"

# 必填标记：字段名后缀 *
_REQUIRED_SUFFIX = "*"

# 类型/示例解析（Field Remark 内）
_TYPE_RE = re.compile(r"Type\s*:\s*([A-Za-z/]+)(?:\((\d+)\))?", re.IGNORECASE)
_EXAMPLE_RE = re.compile(r"Example\s*:\s*([^\n]*)", re.IGNORECASE)
# 场景里的调用行：Web Service Call: POST http://.../api_xxx?key=xxxxx
_CALL_RE = re.compile(
    r"Web\s*Service\s*Call\s*:\s*(GET|POST|PUT|DELETE|PATCH|HEAD|OPTIONS)\s+(\S+)",
    re.IGNORECASE,
)
# 接口标题：Check Mobile Duplicate (api_checkMobileDuplicate)
_TITLE_RE = re.compile(r"\(\s*(api_[A-Za-z0-9_]+)\s*\)")


@dataclass
class CellVal:
    """带坐标的单元格值。"""

    row: int
    col: int
    value: str

    @property
    def coord(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"


@dataclass
class ParsedField:
    """解析出的单个字段（入参或出参）。"""

    name: str
    required: bool
    type: str | None
    length: int | None
    example: str | None
    description: str | None
    level: str | None          # Field層級（父字段，如 Data）
    match: str | None          # 匹配列 √/×/？
    diff: str | None           # 說明/差異
    follow_up: str | None
    questions: str | None
    source: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedInterface:
    """解析出的单个接口 Sheet。"""

    sheet: str
    api_name: str | None = None      # api_xxx
    title: str | None = None         # 场景标题
    scenario: str | None = None      # 场景描述全文
    http_method: str | None = None   # 从调用行解析
    call_url: str | None = None      # 原始调用地址（可能为占位）
    path: str | None = None          # 从调用行提取的 path
    request_fields: list[ParsedField] = field(default_factory=list)
    response_fields: list[ParsedField] = field(default_factory=list)
    error_codes: list[dict[str, Any]] = field(default_factory=list)
    source_rows: list[int] = field(default_factory=list)


class ExcelAPIParser:
    """美心 Excel API 文档解析器。"""

    def load(self, content: bytes) -> "ExcelAPIParser":
        """从字节流加载 workbook（data_only 取公式计算值）。"""
        import io

        self._wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        return self

    @property
    def workbook(self):
        """已加载的 workbook（供调用方按 sheet 名取工作表）。"""
        if getattr(self, "_wb", None) is None:
            raise RuntimeError("请先调用 load() 加载文档")
        return self._wb

    # ---------- Sheet 级 ----------

    def sheet_names(self) -> list[str]:
        return list(self._wb.sheetnames)

    def is_guide_sheet(self, ws: Worksheet) -> bool:
        """判断是否导读 Sheet：含公共申请字段区（model/method/appid/ts/sign）。"""
        text = self._sheet_text(ws).lower()
        return "model" in text and "appid" in text and "sign" in text and "method" in text

    def parse_guide(self, ws: Worksheet) -> dict[str, Any]:
        """解析导读 Sheet -> 公共规则。"""
        rules: dict[str, str] = {}
        sign_text = ""
        for row in ws.iter_rows():
            for c in row:
                v = self._clean(c.value)
                if not v:
                    continue
                low = v.lower()
                if low in {"model", "method", "appid", "ts", "sign"}:
                    # 同行右侧单元格为说明
                    remark = self._clean(ws.cell(row=c.row, column=c.column + 1).value)
                    rules[low] = remark or v
                elif "md5" in low or "签名" in v or "簽名" in v:
                    if len(v) > len(sign_text):
                        sign_text = v

        return {
            "model": rules.get("model"),
            "method_rule": rules.get("method"),
            "appid_rule": rules.get("appid"),
            "ts_rule": rules.get("ts"),
            "sign_rule": rules.get("sign"),
            "sign_text": sign_text or None,
            "sign": self._parse_sign_formula(sign_text),
        }

    def parse_interface_sheet(self, ws: Worksheet) -> ParsedInterface:
        """解析单个接口 Sheet。"""
        cells = self._merged_aware_cells(ws)
        header = self._find_header(cells)
        colmap = header["colmap"] if header else {}

        iface = ParsedInterface(sheet=ws.title)
        section = None  # "input" | "output" | None
        current_level: str | None = None

        # 按行分组遍历（跳过表头行）
        rows = self._group_by_row(cells)
        for row_idx in sorted(rows.keys()):
            if header and row_idx == header["row"]:
                continue
            row_cells = {c.col: c for c in rows[row_idx]}
            col_a = self._cell_at(row_cells, 1)
            a_val = col_a.value if col_a else ""
            field_cell_peek = self._cell_at(row_cells, colmap.get("field", 3))
            has_field_name = bool(field_cell_peek and field_cell_peek.value.strip()
                                  and field_cell_peek.value.strip().lower() not in {"field", "字段"})

            # 场景行（API使用場景）
            if "使用場景" in a_val or "使用场景" in a_val:
                self._fill_scenario(iface, row_cells, ws.title, row_idx)
                continue

            # 分区切换（合并单元格会把分区标记填到每一行，因此分区行可能也是字段行）
            a_low = a_val.lower().replace(" ", "")
            if a_low.startswith(_SECTION_INPUT):
                section = "input"
                if not has_field_name:
                    continue
            elif a_low.startswith(_SECTION_OUTPUT):
                section = "output"
                if not has_field_name:
                    continue
            elif a_low.startswith(_RETURN_CODE_MARK):
                section = "error_code"
                if not has_field_name:
                    continue

            # 错误码区：C 列=码，D 列=含义
            if section == "error_code":
                code = self._cell_at(row_cells, colmap.get("field", 3))
                meaning = self._cell_at(row_cells, colmap.get("field_remark", 4))
                if code and code.value.strip().isdigit():
                    iface.error_codes.append({
                        "code": code.value.strip(),
                        "meaning": meaning.value if meaning else "",
                        "source": {"sheet": ws.title, "row": row_idx},
                    })
                    iface.source_rows.append(row_idx)
                continue

            # 字段行（入参/出参）
            field_cell = field_cell_peek
            if not field_cell or not field_cell.value.strip():
                continue
            name_raw = field_cell.value.strip()
            # 跳过表头残留/分区标题
            if name_raw.lower() in {"field", "字段"}:
                continue
            # 尚未进入任何分区（容错：默认按入参）
            if section is None:
                section = "input"

            level_cell = self._cell_at(row_cells, colmap.get("field_level", 2))
            level_val = level_cell.value if level_cell else None
            # Field層級列出现非 "-" 值时更新当前层级（如 Data）
            if level_val and level_val.strip() not in {"", "-"}:
                current_level = level_val.strip()

            remark_cell = self._cell_at(row_cells, colmap.get("field_remark", 4))
            remark = remark_cell.value if remark_cell else ""

            ftype, flen = self._parse_type(remark)
            example = self._parse_example(remark)
            required = name_raw.endswith(_REQUIRED_SUFFIX)
            name = name_raw[:-1] if required else name_raw

            match_cell = self._cell_at(row_cells, colmap.get("match", 5))
            diff_cell = self._cell_at(row_cells, colmap.get("diff", 6))
            follow_cell = self._cell_at(row_cells, colmap.get("follow_up", 7))
            ques_cell = self._cell_at(row_cells, colmap.get("questions", 8))

            pf = ParsedField(
                name=name,
                required=required,
                type=ftype,
                length=flen,
                example=example,
                description=self._parse_description(remark),
                level=current_level,
                match=match_cell.value if match_cell else None,
                diff=diff_cell.value if diff_cell else None,
                follow_up=follow_cell.value if follow_cell else None,
                questions=ques_cell.value if ques_cell else None,
                source={
                    "sheet": ws.title,
                    "row": row_idx,
                    "col": field_cell.col,
                    "cell": field_cell.coord,
                },
            )
            if section == "input":
                iface.request_fields.append(pf)
            elif section == "output":
                iface.response_fields.append(pf)
            iface.source_rows.append(row_idx)

        return iface

    # ---------- 内部工具 ----------

    def _merged_aware_cells(self, ws: Worksheet) -> list[CellVal]:
        """展开合并单元格，返回所有非空单元格（带坐标）。"""
        merged_map: dict[tuple[int, int], str] = {}
        for rng in ws.merged_cells.ranges:
            top_left = ws.cell(row=rng.min_row, column=rng.min_col).value
            if top_left is None:
                continue
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merged_map[(r, c)] = str(top_left)

        out: list[CellVal] = []
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None and (c.row, c.column) in merged_map:
                    v = merged_map[(c.row, c.column)]
                s = self._clean(v)
                if s:
                    out.append(CellVal(row=c.row, col=c.column, value=s))
        return out

    def _find_header(self, cells: list[CellVal]) -> dict[str, Any] | None:
        """定位表头行（含 Field 且含 Field Remark），返回列映射。"""
        rows = self._group_by_row(cells)
        for row_idx in sorted(rows.keys()):
            colmap: dict[str, int] = {}
            for c in rows[row_idx]:
                key = _HEADER_ALIASES.get(c.value.lower())
                if key:
                    colmap[key] = c.col
            if "field" in colmap and "field_remark" in colmap:
                return {"row": row_idx, "colmap": colmap}
        return None

    def _fill_scenario(self, iface: ParsedInterface, row_cells: dict[int, CellVal], sheet: str, row_idx: int) -> None:
        # 场景标题在 C 列，描述在 D 列
        title_cell = self._cell_at(row_cells, 3)
        desc_cell = self._cell_at(row_cells, 4)
        iface.title = title_cell.value if title_cell else None
        desc = desc_cell.value if desc_cell else ""
        iface.scenario = desc or None
        m = _TITLE_RE.search(iface.title or "")
        if m:
            iface.api_name = m.group(1)
        call = _CALL_RE.search(desc or "")
        if call:
            iface.http_method = call.group(1).upper()
            iface.call_url = call.group(2)
            iface.path = self._path_from_url(call.group(2))
        iface.source_rows.append(row_idx)

    @staticmethod
    def _path_from_url(url: str) -> str | None:
        marker = url.find("/api/")
        return url[marker:] if marker >= 0 else None

    @staticmethod
    def _parse_type(remark: str) -> tuple[str | None, int | None]:
        m = _TYPE_RE.search(remark or "")
        if not m:
            return None, None
        return m.group(1), int(m.group(2)) if m.group(2) else None

    @staticmethod
    def _parse_example(remark: str) -> str | None:
        m = _EXAMPLE_RE.search(remark or "")
        if not m:
            return None
        val = m.group(1).strip()
        return val if val else None

    @staticmethod
    def _parse_description(remark: str) -> str | None:
        """去掉 Type/Example 行后的描述文本。"""
        if not remark:
            return None
        lines = []
        for line in remark.splitlines():
            s = line.strip()
            if not s:
                continue
            if _TYPE_RE.match(s) or _EXAMPLE_RE.match(s):
                continue
            lines.append(s)
        return "\n".join(lines) or None

    def _parse_sign_formula(self, text: str) -> dict[str, Any] | None:
        """从签名说明文本解析算法与字段拼接顺序。"""
        if not text:
            return None
        low = text.lower()
        if "md5" not in low:
            return None
        # 识别拼接字段顺序：appid + appsercret + ts + msg
        order = []
        for token in ("appid", "appsercret", "appsecret", "ts", "msg"):
            if token in low:
                order.append("appsercret" if token == "appsecret" else token)
        return {
            "algo": "md5",
            "case": "lower",
            "digest_length": 32,
            "concat_order": order or ["appid", "appsercret", "ts", "msg"],
            "raw": text,
        }

    def _sheet_text(self, ws: Worksheet) -> str:
        parts = []
        for row in ws.iter_rows():
            for c in row:
                v = self._clean(c.value)
                if v:
                    parts.append(v)
        return "\n".join(parts)

    @staticmethod
    def _group_by_row(cells: list[CellVal]) -> dict[int, list[CellVal]]:
        out: dict[int, list[CellVal]] = {}
        for c in cells:
            out.setdefault(c.row, []).append(c)
        return out

    @staticmethod
    def _cell_at(row_cells: dict[int, CellVal], col: int | None) -> CellVal | None:
        if col is None:
            return None
        return row_cells.get(col)

    @staticmethod
    def _clean(v: Any) -> str:
        if v is None:
            return ""
        return str(v).replace("\r\n", "\n").replace("\r", "\n").strip()
