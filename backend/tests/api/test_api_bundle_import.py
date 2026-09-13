# -*- coding: utf-8 -*-
"""统一接口资产包 API 测试（parse-bundle / confirm / 幂等 / 冲突 / 回滚）。

覆盖可交付项：
1. parse-bundle 保存 Bundle / Source / 解析结果 / 匹配结果；
2. confirm 在单事务内创建 Folder、APIEndpoint、变量、断言、依赖；
3. 幂等键 = 项目 + method + normalized_path；
4. 冲突确认校验（high 未确认 → 拒绝导入）；
5. 失败回滚（不留半成品）。

测试策略：
- 使用「内联最小样本」保证可重复（真实美心三文件 HTTP 验收见
  ``backend/scripts/mx_bundle_acceptance.py``）；
- 通过真实 HTTP（httpx ASGITransport）调用路由；
- 使用专用 PostgreSQL 测试库（模型使用 JSONB/UUID，且回滚语义依赖真实事务）。
"""

from __future__ import annotations

import json
import uuid
from pathlib import Path

import pytest


# ---------------------------------------------------------------------------
# 样本构造
# ---------------------------------------------------------------------------

def _excel_bytes() -> bytes:
    """生成符合美心模板结构的最小 xlsx（1 导读 + 2 接口）。

    布局严格对齐真实文件 ``YX CRM API Spec For Non POS``：
    - 第 1 行表头：A=入參 & 出參 / B=Field層級 / C=Field / D=Field Remark
      / E=匹配 / F=說明/差異 / G=Follow up / H=疑問清單
    - 第 2 行场景：A=API使用場景，C=标题（含 (api_xxx)），D=含 Web Service Call 的描述
    - 第 3 行起：A 列（合并）为分区标记 ``Field（入參）`` / ``Field（出參）`` / ``Return Code``，
      C 列字段名（``*`` 结尾表示必填），D 列 ``Type: xxx(n) Example: yyy``
    """
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    guide = wb.create_sheet("Guide")
    for row in (
        ["model", "api"],
        ["method", "业务方法名"],
        ["appid", "应用ID"],
        ["ts", "时间戳"],
        ["sign", "MD5(appid+appsercret+ts+msg)"],
    ):
        guide.append(row)

    for api_name, method, path, title in (
        ("api_checkMobileDuplicate", "POST", "/api/cop/json/api_checkMobileDuplicate", "Check Mobile Duplicate"),
        ("api_getMemberAct", "POST", "/api/cop/json/api_getMemberAct", "Get Member Act"),
    ):
        ws = wb.create_sheet(f"COP {api_name}"[:31])
        ws.append(["入參 & 出參", "Field層級", "Field", "Field Remark", "匹配", "說明/差異", "Follow up", "疑問清單"])
        ws.append([
            "API使用場景",
            "-",
            f"{title} ({api_name})",
            f"Source: Eatizen One App\nWeb Service Call: {method} {path}",
        ])
        # 入参区
        ws.append(["Field（入參）", "-", "MobilePhone*", "Type: Text(20)\nExample: 61234567", "√"])
        ws.append([None, "-", "CountryCode", "Type: Text(5)\nExample: 852", "√"])
        # 出参区
        ws.append(["Field（出參）", "-", "Existed", "Type: Text(10)\nExample: MX1", "√"])
        # 返回码区：分区标记必须在 A 列（与真实文件一致，解析器只读 A 列判定分区），
        # C 列=码，D 列=含义
        ws.append(["Return Code", "-", "0", "成功"])
        ws.append([None, "-", "1", "参数错误"])

        # 合并分区标记列（与真实文件一致）
        ws.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=1)
        ws.merge_cells(start_row=6, start_column=1, end_row=7, end_column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


COLLECTION = {
    "info": {
        "name": "美心测试集合",
        "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json",
    },
    "item": [
        {
            "name": "api_checkMobileDuplicate",
            "request": {
                "method": "POST",
                "header": [{"key": "method", "value": "checkMobileDuplicate"}],
                "url": {"raw": "{{url}}/cop/json", "host": ["{{url}}"], "path": ["cop", "json"]},
                "body": {
                    "mode": "raw",
                    "raw": json.dumps({"MobilePhone": "61234567", "CountryCode": "852"}),
                },
            },
            "event": [{
                "listen": "test",
                "script": {"exec": [
                    "pm.test('status 200', function () { pm.response.to.have.status(200); });"
                ]},
            }],
        },
        {
            "name": "api_getMemberAct",
            "request": {
                "method": "POST",
                "header": [{"key": "method", "value": "getMemberAct"}],
                "url": {"raw": "{{url}}/cop/json", "host": ["{{url}}"], "path": ["cop", "json"]},
                "body": {
                    "mode": "raw",
                    "raw": json.dumps({"memberID": "{{memberID}}", "sessionKey": "{{sessionKey}}"}),
                },
            },
        },
    ],
    "variable": [{"key": "url", "value": "https://uat.example.com"}],
}

ENVIRONMENT = {
    "id": "env-1",
    "name": "UAT eatizen",
    "values": [
        {"key": "url", "value": "https://uat.example.com", "type": "default", "enabled": True},
        {"key": "api-appid", "value": "app-1", "type": "default", "enabled": True},
        {"key": "api-secret", "value": "s3cr3t", "type": "secret", "enabled": True},
    ],
    "_postman_variable_scope": "environment",
}


# ---------------------------------------------------------------------------
# 夹具：专用 PostgreSQL 测试库
# ---------------------------------------------------------------------------

@pytest.fixture()
async def client():
    """仅挂载 api_bundles 路由的应用 + 专用 PostgreSQL 测试库。

    engine 必须在测试事件循环内创建，否则 asyncpg 连接池会绑定到已关闭的循环。
    """
    import os

    import httpx
    from fastapi import FastAPI
    from sqlalchemy import create_engine, text
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from app.api.deps import get_current_user_id, get_db
    from app.api.v2.api_bundles import router
    from app.config.settings import settings
    from app.models.base import Base

    # 注册全部模型
    from app.models import api_asset as _api_asset  # noqa: F401
    from app.models import api_doc_parse_task as _task  # noqa: F401
    from app.models import api_endpoint as _api_endpoint  # noqa: F401
    from app.models import attachment as _attachment  # noqa: F401
    from app.models import folder as _folder  # noqa: F401
    from app.models import project as _project  # noqa: F401
    from app.models import user as _user  # noqa: F401

    if os.getenv("PGPASSWORD") is None:
        os.environ["PGPASSWORD"] = "postgres"

    base_url = settings.postgres_sync_url
    test_name = "bundle_test_db"
    sync_url = base_url.rsplit("/", 1)[0] + "/" + test_name
    async_url = sync_url.replace("postgresql://", "postgresql+asyncpg://")
    admin_url = base_url.rsplit("/", 1)[0] + "/postgres"

    def _admin_exec(sql: str) -> None:
        eng = create_engine(admin_url, isolation_level="AUTOCOMMIT")
        with eng.connect() as conn:
            conn.execute(text(sql))
        eng.dispose()

    _admin_exec(f'DROP DATABASE IF EXISTS "{test_name}" WITH (FORCE)')
    _admin_exec(f'CREATE DATABASE "{test_name}"')

    engine = create_async_engine(async_url, pool_pre_ping=True)
    session_factory = async_sessionmaker(engine, expire_on_commit=False)

    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async with session_factory() as s:
        from app.models.project import Project
        from app.models.user import User

        user = User(
            id=uuid.uuid4(),
            email="bundle-test@example.com",
            username="bundle-test",
            password_hash="x",
            is_active=True,
        )
        s.add(user)
        await s.flush()
        s.add(Project(identifier="PR-TEST", name="统一导入测试项目", created_by=user.id))
        await s.commit()

    app = FastAPI()
    app.include_router(router, prefix="/api/v2")
    app.dependency_overrides[get_current_user_id] = lambda: "user-1"

    async def _db_dep():
        async with session_factory() as s:
            yield s

    app.dependency_overrides[get_db] = _db_dep

    transport = httpx.ASGITransport(app=app)
    api = httpx.AsyncClient(transport=transport, base_url="http://test")

    try:
        yield api
    finally:
        await api.aclose()
        await engine.dispose()
        _admin_exec(f'DROP DATABASE IF EXISTS "{test_name}" WITH (FORCE)')


BASE = "/api/v2/projects/PR-TEST/api-bundles"


async def parse(client, bundle_name: str, *, excel=True, collection=True, environment=True):
    """调用 parse-bundle（文本字段走 data=，文件字段走 files=）。"""
    data = {"bundle_name": bundle_name, "system_type": "crm_non_pos"}
    files = {}
    if excel:
        files["excel_file"] = (
            "spec.xlsx", _excel_bytes(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    if collection:
        files["collection_file"] = (
            "collection.json", json.dumps(COLLECTION).encode("utf-8"), "application/json"
        )
    if environment:
        files["environment_file"] = (
            "env.json", json.dumps(ENVIRONMENT).encode("utf-8"), "application/json"
        )
    return await client.post(f"{BASE}/parse-bundle", data=data, files=files)



def _unconfirmed(blocked) -> list:
    """从 409 响应取未确认冲突清单（兼容 detail 与 data 两种承载）。"""
    body = blocked.json()
    for key in ("detail", "data"):
        payload = body.get(key)
        if isinstance(payload, dict) and payload.get("unconfirmed_conflicts"):
            return payload["unconfirmed_conflicts"]
    return []


async def confirm_all(client, bundle_id: str):
    """确认全部 high 冲突后执行 confirm，返回最终 confirm 响应。"""
    blocked = await client.post(f"{BASE}/{bundle_id}/confirm", json={})
    if blocked.status_code == 409:
        conflicts = _unconfirmed(blocked)
        if not conflicts:
            raise AssertionError(f"409 响应缺少未确认冲突清单: {blocked.text[:300]}")
        resolved = await client.post(
            f"{BASE}/{bundle_id}/conflicts/confirm",
            json={"resolutions": [
                {"conflict_id": c["id"], "resolution": "accept", "note": "测试确认"}
                for c in conflicts
            ]},
        )
        if resolved.status_code != 200:
            raise AssertionError(
                f"冲突确认失败 {resolved.status_code}: {resolved.text[:300]}"
            )
    elif blocked.status_code not in (200,):
        raise AssertionError(f"confirm 返回异常 {blocked.status_code}: {blocked.text[:300]}")

    final = await client.post(f"{BASE}/{bundle_id}/confirm", json={})
    if final.status_code != 200:
        raise AssertionError(f"confirm 最终失败 {final.status_code}: {final.text[:500]}")
    return final


# ---------------------------------------------------------------------------
# 1. parse-bundle
# ---------------------------------------------------------------------------

class TestParseBundle:
    """parse-bundle 保存 Bundle / Source / 解析结果 / 匹配结果"""

    async def test_saves_bundle_sources_and_contracts(self, client):
        resp = await parse(client, "bundle-A")
        assert resp.status_code == 200, resp.text
        data = resp.json()["data"]

        assert data["bundle_id"]
        assert data["status"] == "parsed"
        assert len(data["sources"]) == 3, "三个来源都要落库"
        assert {s["source_type"] for s in data["sources"]} == {
            "excel_mx", "postman_collection", "postman_environment"
        }
        assert data["stats"]["excel_interfaces"] == 2
        assert data["stats"]["collection_requests"] == 2
        assert data["stats"]["environment_variables"] == 3

        assert len(data["contracts"]) >= 2, "匹配结果（契约）要落库"
        paths = {c["normalized_path"] for c in data["contracts"]}
        assert any("checkmobileduplicate" in p for p in paths)

    async def test_is_idempotent(self, client):
        first = await parse(client, "bundle-B")
        second = await parse(client, "bundle-B")
        assert first.status_code == 200 and second.status_code == 200
        d1, d2 = first.json()["data"], second.json()["data"]

        assert d1["bundle_id"] == d2["bundle_id"], "同名 bundle 应复用"
        assert len(d2["sources"]) == len(d1["sources"]), "来源不重复登记"
        assert len(d2["contracts"]) == len(d1["contracts"]), "契约不重复登记"

    async def test_rejects_unknown_source_type(self, client):
        resp = await client.post(
            f"{BASE}/parse-bundle",
            data={"bundle_name": "bundle-C"},
            files={"unknown_file": ("x.txt", b"x", "text/plain")},
        )
        assert resp.status_code == 400, resp.text

    async def test_requires_at_least_one_source(self, client):
        resp = await client.post(f"{BASE}/parse-bundle", data={"bundle_name": "bundle-D"})
        assert resp.status_code == 400

    async def test_variables_saved_without_plaintext(self, client):
        parsed = await parse(client, "bundle-E")
        bundle_id = parsed.json()["data"]["bundle_id"]
        resp = await client.get(f"{BASE}/{bundle_id}/variables")

        assert resp.status_code == 200
        variables = resp.json()["data"]
        assert len(variables) == 3
        for v in variables:
            assert v["value_ref"].startswith("env:"), "变量只能存引用，不能存明文"
        secret = next(v for v in variables if v["name"] == "api-secret")
        assert secret["is_secret"] is True
        assert "s3cr3t" not in json.dumps(variables, ensure_ascii=False)


# ---------------------------------------------------------------------------
# 2. confirm：单事务创建
# ---------------------------------------------------------------------------

class TestConfirmCreatesAssets:
    """confirm 在事务内创建 Folder / APIEndpoint / 变量 / 断言 / 依赖"""

    async def test_creates_folder_endpoints_and_assertions(self, client):
        parsed = await parse(client, "confirm-A")
        bundle_id = parsed.json()["data"]["bundle_id"]

        resp = await confirm_all(client, bundle_id)
        assert resp.status_code == 200, resp.text
        result = resp.json()["data"]
        assert result["imported_count"] > 0
        assert result["folder_ids"], "必须创建 Folder"
        assert not result["failed"]

        assets = await client.get(f"{BASE}/{bundle_id}/assets")
        assert assets.status_code == 200
        stats = assets.json()["data"]
        assert stats["contracts"] > 0
        assert stats["contracts_with_endpoint"] == stats["contracts"], "每个契约都要回填端点"
        assert stats["request_templates"] > 0, "请求模板要落库"
        assert stats["assertions"] > 0, "断言要落库"

    async def test_is_idempotent_by_project_method_path(self, client):
        """幂等键 = 项目 + method + normalized_path"""
        parsed = await parse(client, "confirm-B")
        bundle_id = parsed.json()["data"]["bundle_id"]

        first = await confirm_all(client, bundle_id)
        second = await client.post(f"{BASE}/{bundle_id}/confirm", json={})
        assets = await client.get(f"{BASE}/{bundle_id}/assets")

        d1, d2 = first.json()["data"], second.json()["data"]
        assert d1["imported_count"] > 0
        assert d2["imported_count"] == 0, "重复 confirm 不应新建端点"
        assert d2["skipped_count"] == d1["imported_count"] + d1["skipped_count"]
        assert assets.json()["data"]["contracts_with_endpoint"] == d1["imported_count"]

    async def test_reuses_endpoint_across_bundles(self, client):
        """不同 bundle 的同一 method+path 复用同一个端点"""
        ids = []
        for name in ("cross-1", "cross-2"):
            parsed = await parse(client, name)
            bid = parsed.json()["data"]["bundle_id"]
            ids.append(bid)
            await confirm_all(client, bid)

        a = await client.get(f"{BASE}/{ids[0]}")
        b = await client.get(f"{BASE}/{ids[1]}")
        eps_a = {c["normalized_path"]: c["imported_endpoint_id"]
                 for c in a.json()["data"]["contracts"]}
        eps_b = {c["normalized_path"]: c["imported_endpoint_id"]
                 for c in b.json()["data"]["contracts"]}
        common = set(eps_a) & set(eps_b)
        assert common, "两个 bundle 应有相同接口"
        for path in common:
            assert eps_a[path] == eps_b[path], f"同一 method+path 必须复用同一端点: {path}"

    async def test_excel_error_codes_become_assertions(self, client):
        """样本含 Return Code 区 → confirm 后应落库 excel_contract 断言。

        回归：返回码区曾因分区标记与去空格文本不匹配而永不解析，
        断言恒为 0，且资产统计无法区分断言来源。
        """
        parsed = await parse(client, "excel-assertions")
        bundle_id = parsed.json()["data"]["bundle_id"]
        ok = await confirm_all(client, bundle_id)
        assert ok.status_code == 200, ok.text

        assets = (await client.get(f"{BASE}/{bundle_id}/assets")).json()["data"]
        by_source = assets.get("assertions_by_source") or {}
        assert int(by_source.get("excel_contract", 0)) > 0, (
            f"Excel 返回码应产生 excel_contract 断言，实际 {by_source}"
        )
        assert assets["assertions"] == sum(int(v) for v in by_source.values())



# ---------------------------------------------------------------------------
# 3. 冲突确认校验
# ---------------------------------------------------------------------------

class TestConflictGating:
    """存在未确认 high 冲突时必须拒绝导入"""

    async def test_high_conflict_blocks_confirm(self, client):
        parsed = await parse(client, "gate-A")
        bundle_id = parsed.json()["data"]["bundle_id"]
        resp = await client.post(f"{BASE}/{bundle_id}/confirm", json={})

        assert resp.status_code == 409, resp.text
        body = resp.json()
        # 全局异常处理器把 HTTPException.detail 放在 detail 键下
        payload = body.get("detail") or body.get("data") or {}
        assert payload.get("unconfirmed_count", 0) > 0, body
        assert payload["unconfirmed_conflicts"][0]["conflict_type"] == "env_missing"
        # 被拒绝时不得创建任何端点
        assets = await client.get(f"{BASE}/{bundle_id}/assets")
        assert assets.json()["data"]["contracts_with_endpoint"] == 0

    async def test_confirming_conflicts_unblocks(self, client):
        parsed = await parse(client, "gate-B")
        bundle_id = parsed.json()["data"]["bundle_id"]
        blocked = await client.post(f"{BASE}/{bundle_id}/confirm", json={})
        conflicts = _unconfirmed(blocked)
        resolved = await client.post(
            f"{BASE}/{bundle_id}/conflicts/confirm",
            json={"resolutions": [{"conflict_id": c["id"], "resolution": "accept"}
                                  for c in conflicts]},
        )
        allowed = await client.post(f"{BASE}/{bundle_id}/confirm", json={})

        assert resolved.status_code == 200
        remaining = resolved.json()["data"]["remaining_open"]
        assert not [c for c in remaining if c["severity"] == "high"]
        assert allowed.status_code == 200

    async def test_waive_also_unblocks(self, client):
        parsed = await parse(client, "gate-C")
        bundle_id = parsed.json()["data"]["bundle_id"]
        blocked = await client.post(f"{BASE}/{bundle_id}/confirm", json={})
        conflicts = _unconfirmed(blocked)
        await client.post(
            f"{BASE}/{bundle_id}/conflicts/confirm",
            json={"resolutions": [{"conflict_id": c["id"], "resolution": "waive"}
                                  for c in conflicts]},
        )
        allowed = await client.post(f"{BASE}/{bundle_id}/confirm", json={})
        assert allowed.status_code == 200

    async def test_invalid_resolution_rejected(self, client):
        parsed = await parse(client, "gate-D")
        bundle_id = parsed.json()["data"]["bundle_id"]
        blocked = await client.post(f"{BASE}/{bundle_id}/confirm", json={})
        conflicts = _unconfirmed(blocked)
        resp = await client.post(
            f"{BASE}/{bundle_id}/conflicts/confirm",
            json={"resolutions": [{"conflict_id": conflicts[0]["id"], "resolution": "maybe"}]},
        )
        assert resp.status_code == 400


# ---------------------------------------------------------------------------
# 4. 失败回滚
# ---------------------------------------------------------------------------

class TestRollback:
    """失败必须整体回滚，不留半成品"""

    async def test_unknown_contract_id_rolls_back(self, client):
        parsed = await parse(client, "rb-A")
        bundle_id = parsed.json()["data"]["bundle_id"]
        blocked = await client.post(f"{BASE}/{bundle_id}/confirm", json={})
        conflicts = _unconfirmed(blocked)
        await client.post(
            f"{BASE}/{bundle_id}/conflicts/confirm",
            json={"resolutions": [{"conflict_id": c["id"], "resolution": "accept"}
                                  for c in conflicts]},
        )
        before = await client.get(f"{BASE}/{bundle_id}/assets")
        resp = await client.post(
            f"{BASE}/{bundle_id}/confirm", json={"contract_ids": [str(uuid.uuid4())]}
        )
        after = await client.get(f"{BASE}/{bundle_id}/assets")

        assert resp.status_code == 400, resp.text
        assert "契约" in resp.text
        assert after.json()["data"] == before.json()["data"], "失败后不应留下半成品"

    async def test_unknown_bundle_returns_404(self, client):
        resp = await client.post(f"{BASE}/{uuid.uuid4()}/confirm", json={})
        assert resp.status_code == 404

    async def test_get_unknown_bundle_returns_404(self, client):
        resp = await client.get(f"{BASE}/{uuid.uuid4()}")
        assert resp.status_code == 404


# ---------------------------------------------------------------------------
# 5. 归一化路径（幂等键稳定性）
# ---------------------------------------------------------------------------

class TestNormalizationForIdempotency:
    """归一化路径规则必须稳定，否则幂等键会漂移"""

    @pytest.mark.parametrize(
        "raw,expected",
        [
            ("https://uat.maxims.com/api/json/memLogin", "/api/json/memlogin"),
            ("{{url}}/api/json/memLogin", "/api/json/memlogin"),
            ("/api/json/memLogin/", "/api/json/memlogin"),
            ("/api/json/MemLogin", "/api/json/memlogin"),
            ("/api/json/checkMobile?country=852", "/api/json/checkmobile"),
            ("/api/icdp//v1/cs/cases/page", "/api/icdp/v1/cs/cases/page"),
            ("/api/json/api_memLogin", "/api/json/memlogin"),
            ("", ""),
        ],
    )
    def test_normalize_import_path(self, raw, expected):
        from app.services.api_bundle_service import normalize_import_path

        assert normalize_import_path(raw) == expected


# ---------------------------------------------------------------------------
# 6. 结构就绪校验（P0-4 迁移漂移防护）
# ---------------------------------------------------------------------------

def _schema_test_ddl(sql: str) -> None:
    """在测试库上执行 DDL（绕过 ORM 缓存，直接改真实约束）。"""
    from sqlalchemy import create_engine, text

    from app.config.settings import settings

    url = settings.postgres_sync_url.rsplit("/", 1)[0] + "/bundle_test_db"
    engine = create_engine(url, isolation_level="AUTOCOMMIT")
    try:
        with engine.connect() as conn:
            conn.execute(text(sql))
    finally:
        engine.dispose()


class TestSchemaReadinessGuard:
    """幂等键依赖 uq_endpoint_project_method_normpath 唯一约束。

    约束缺失时必须**硬拒绝导入**（503），而不是降级成"重复创建端点"。
    """

    async def test_confirm_503_when_idempotency_constraint_missing(self, client):
        parsed = await parse(client, "schema-guard")
        bundle_id = parsed.json()["data"]["bundle_id"]

        _schema_test_ddl(
            "ALTER TABLE api_endpoints "
            "DROP CONSTRAINT IF EXISTS uq_endpoint_project_method_normpath"
        )
        try:
            resp = await client.post(f"{BASE}/{bundle_id}/confirm", json={})
            assert resp.status_code == 503, resp.text
            body = resp.json()
            detail = body.get("detail") or body
            assert detail.get("error") == "schema_not_ready"
            assert any("唯一约束缺失" in m for m in detail.get("missing", []))
            assert detail.get("remedy") == "alembic upgrade head"
        finally:
            _schema_test_ddl(
                "ALTER TABLE api_endpoints ADD CONSTRAINT "
                "uq_endpoint_project_method_normpath "
                "UNIQUE (project_id, method, normalized_path)"
            )

        # 约束恢复后必须可正常导入（证明拒绝只源于约束缺失）。
        # 注意：样本本身含 high 冲突，需先确认，否则会被冲突门禁拦成 409。
        ok = await confirm_all(client, bundle_id)
        assert ok.status_code == 200, ok.text
        assert ok.json()["data"]["imported_count"] > 0
        assert ok.json()["data"]["schema_check"] == "passed"

    async def test_confirm_does_not_create_endpoints_when_guard_trips(self, client):
        """守卫触发时不得留下任何端点（失败必须是原子的）。"""
        parsed = await parse(client, "schema-guard-2")
        bundle_id = parsed.json()["data"]["bundle_id"]

        _schema_test_ddl(
            "ALTER TABLE api_endpoints "
            "DROP CONSTRAINT IF EXISTS uq_endpoint_project_method_normpath"
        )
        try:
            resp = await client.post(f"{BASE}/{bundle_id}/confirm", json={})
            assert resp.status_code == 503, resp.text
        finally:
            _schema_test_ddl(
                "ALTER TABLE api_endpoints ADD CONSTRAINT "
                "uq_endpoint_project_method_normpath "
                "UNIQUE (project_id, method, normalized_path)"
            )

        assets = await client.get(f"{BASE}/{bundle_id}/assets")
        assert assets.status_code == 200, assets.text
        data = assets.json()["data"]
        assert data["contracts_with_endpoint"] == 0, "守卫触发时不应回填任何端点"
        assert data["request_templates"] == 0

    async def test_idempotency_helpers(self):
        """约束覆盖判定：顺序不敏感、大小写不敏感、缺列则判缺失。"""
        from app.services.api_bundle_service import _constraint_covers_key

        assert _constraint_covers_key(["project_id", "method", "normalized_path"])
        assert _constraint_covers_key(["METHOD", "PROJECT_ID", "NORMALIZED_PATH"])
        assert _constraint_covers_key(["method", "normalized_path", "project_id", "folder_id"])
        assert not _constraint_covers_key(["project_id", "method"])
        assert not _constraint_covers_key([])

    async def test_ensure_schema_ready_passes_on_current_db(self, client):
        """当前测试库（含约束）必须通过就绪校验。"""
        from app.api.deps import get_db
        from app.services.api_bundle_service import ApiBundleService

        dep = client._transport.app.dependency_overrides[get_db]
        async for session in dep():
            await ApiBundleService(session).ensure_schema_ready()
            break
