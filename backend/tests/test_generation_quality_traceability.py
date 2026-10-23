# -*- coding: utf-8 -*-
"""本轮质量问题收口回归：同名标题 / 兜底模块归位 / 待确认项分类。

背景（2026-10-14 复验）：YL-主数据管理 生成结果里
1) 同一字段名出现在「新建/编辑表单字段说明」与「列表页筛选查询区」两处，
   生成出同名用例且无法只看标题判断该测哪个页面；
2) 20 条用例堆在兜底模块（文档名）里——其中有文档归属证据的没归位，
   通用建议又与业务用例混在一起；
3) 待确认项里「文档歧义」与「人工补充」没区分。

这里把三类问题的判定固化成回归，避免后续改动回退。
"""

from __future__ import annotations

import io

from app.api.v2.ai import (
    AMBIGUITY_DOCUMENT,
    AMBIGUITY_MANUAL_SUPPLEMENT,
    GENERIC_SUGGESTION_MODULE,
    _build_minimal_cases,
    _build_import_supplement_cases,
    _resolve_duplicate_titles,
    _case_scenario_label,
)
from app.schemas.test_case import TestCaseCreate, TestStepCreate
from app.services.requirement_structure import (
    locate_evidence,
    locate_table_locations,
    parse_requirement_document,
)


def _build_docx(paragraphs: list[str], tables: list[list[list[str]]]) -> bytes:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    for rows in tables:
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


#: 新建/编辑表单「字段说明」表（同一字段在此处是表单字段）
FORM_TABLE = [
    ["字段", "字段类型", "是否必填", "字段说明"],
    ["店鋪編號", "輸入框", "否", "小於等於30個字符；字符類型：大小寫字母、數字"],
    ["所屬品牌", "下拉單選", "是", "選項範圍：品牌主數據中的所有啟用的品牌；"],
]

#: 列表页「筛选查询区」表（同名字段在此处是查询条件，与表单字段不是同一条规则）
FILTER_TABLE = [
    ["字段", "字段类型", "字段校验", "筛选规则"],
    ["店鋪編號", "輸入框", "小於等於30個字符；默認為空，為空時查詢全部", "精準搜索"],
    ["所屬品牌", "下拉單選", "選項範圍：全部、部分", "根據選項篩選"],
]


class TestEvidenceLocator:
    """文档定位：证据 → 行号 → 最近标题 → 一级章节。"""

    TEXT = "\n".join([
        "1. 店鋪主數據管理",
        "1.2.1 新建/編輯店鋪頁",
        "1.2.1.2 字段說明",
        "店鋪編號",
        "小於等於30個字符；字符類型：大小寫字母、數字",
        "1.2.2 門店列表頁",
        "1.2.2.4 篩選查詢區",
        "店鋪編號",
        "小於等於30個字符；默認為空，為空時查詢全部",
        "2. 導入任務",
        "2.1 導入任務列表頁",
    ])
    #: 第 1 页 = 前 5 行（字段说明），第 2 页 = 其余（筛选查询区）
    PAGE_LINE_COUNTS = [5, 6]

    def test_locate_returns_line_heading_and_section(self):
        found = locate_evidence(self.TEXT, "小於等於30個字符；默認為空，為空時查詢全部")
        assert found is not None
        assert found["line"] == 8
        assert found["heading"] == "1.2.2.4 篩選查詢區"
        assert found["section"] == "店鋪主數據管理"
        assert found["unique"] is True

    def test_page_restriction_disambiguates_repeated_text(self):
        # 同一段文字在文档里出现两次（跨页各一处）
        repeated = "小於等於30個字符"
        any_hit = locate_evidence(self.TEXT, repeated)
        assert any_hit is not None and any_hit["unique"] is False
        # 限定在第 2 页（筛选查询区所在页）→ 唯一命中筛选查询区
        paged = locate_evidence(
            self.TEXT, repeated, page=2, page_line_counts=self.PAGE_LINE_COUNTS
        )
        assert paged is not None
        assert paged["unique"] is True
        assert paged["heading"] == "1.2.2.4 篩選查詢區"
        assert paged["scope"] == "page"
        # 限定在第 1 页 → 命中字段说明
        paged2 = locate_evidence(
            self.TEXT, repeated, page=1, page_line_counts=self.PAGE_LINE_COUNTS
        )
        assert paged2 is not None and paged2["unique"] is True
        assert paged2["heading"] == "1.2.1.2 字段說明"

    def test_unknown_evidence_returns_none(self):
        assert locate_evidence(self.TEXT, "本文档里根本不存在的句子内容") is None
        assert locate_evidence(self.TEXT, "短") is None

    def test_table_locations_use_header_row(self):
        tables = [
            [["字段", "字段类型", "是否必填", "字段说明"], ["店鋪編號", "輸入框", "否", "小於等於30個字符"]],
            [["字段", "字段类型", "字段校验", "筛选规则"], ["店鋪編號", "輸入框", "默認為空", "精準搜索"]],
        ]
        # 表头行原文不在正文里时，退化用该表最长单元格定位（仍是文档原文）
        locations = locate_table_locations(tables, [None, None], self.TEXT)
        assert locations, "至少应定位到可定位的表"
        assert locations[0]["heading"] == "1.2.1.2 字段說明"
        assert locations[0]["section"] == "店鋪主數據管理"

        # 正文按文档顺序包含表内容时（docx 拍平后就是这个形态）→ 两张表各自归位
        interleaved = "\n".join([
            "1. 店鋪主數據管理",
            "1.2.1 新建/編輯店鋪頁",
            "1.2.1.2 字段說明",
            "字段\t字段类型\t是否必填\t字段说明",
            "店鋪編號\t輸入框\t否\t小於等於30個字符；字符類型：大小寫字母、數字",
            "1.2.2 門店列表頁",
            "1.2.2.4 篩選查詢區",
            "字段\t字段类型\t字段校验\t筛选规则",
            "店鋪編號\t輸入框\t小於等於30個字符；默認為空，為空時查詢全部\t精準搜索",
        ])
        located = locate_table_locations(tables, [None, None], interleaved)
        assert located[0]["heading"] == "1.2.1.2 字段說明"
        assert located[1]["heading"] == "1.2.2.4 篩選查詢區"
        assert located[1]["section"] == "店鋪主數據管理"


class TestSameNameDifferentScenario:
    """同一字段名在两处出现：标题必须能区分，且规则追踪不能丢。"""

    def _cases(self):
        content = _build_docx(
            ["1. 店鋪主數據管理", "1.2.1 新建/編輯店鋪頁", "1.2.1.2 字段說明",
             "1.2.2 門店列表頁", "1.2.2.4 篩選查詢區"],
            [FORM_TABLE, FILTER_TABLE],
        )
        structure = parse_requirement_document("YL-主数据管理.docx", None, content)
        cases, gaps_gaps, _gapsskips = _build_minimal_cases(
            structure.full_text, "YL-主数据管理.docx", structure=structure)
        return cases, structure

    def test_no_duplicate_titles_and_rule_ids_preserved(self):
        cases, _structure = self._cases()
        titles = [c.name for c in cases]
        assert len(titles) == len(set(titles)), f"仍有同名标题：{titles}"

        shop_no_cases = [c for c in cases if "店鋪編號" in c.name]
        shop_rules = {c.custom_fields["requirement_rule_id"] for c in shop_no_cases}
        # 表单规则与筛选规则是两条独立规则 → 规则 ID 必须不同且都保留
        assert len(shop_no_cases) >= 2
        assert len(shop_rules) == len(shop_no_cases)

    def test_scenario_labels_come_from_document_headings(self):
        """同名字段用例的标签必须点名**文档里的位置**（子章节/表/行），不能靠猜。"""
        cases, _structure = self._cases()
        labeled = [c for c in cases if c.custom_fields.get("title_disambiguation")]
        assert labeled, "同名字段用例必须带场景标签"
        for case in labeled:
            label = case.custom_fields["title_disambiguation"]
            source = (case.custom_fields["requirement_rule"] or {}).get("source") or {}
            # 标签来自：子章节标题 / 表号行号 / 规则 ID —— 三者都是可追溯的原文位置
            assert any(marker in label for marker in ("字段說明", "字段说明", "篩選查詢",
                                                      "筛选查询", "來源", "来源")
                       or source.get("table") is not None
                       or case.custom_fields["requirement_rule_id"] in label), label

    def test_attributed_by_document_evidence_not_guess(self):
        cases, _structure = self._cases()
        business = [c for c in cases if not c.custom_fields.get("companion_template")]
        assert business
        for case in business:
            assert case.module == "店鋪主數據管理", case.module
            assert case.custom_fields.get("module_source") in {
                "prd_section", "prd_section_evidence", "table_header"}


class TestCompanionCasesAreNotForceClassified:
    def test_companions_get_generic_module_and_tag(self):
        content = _build_docx(
            ["1. 店鋪主數據管理", "1.2.1.2 字段說明",
             "維護店鋪的主數據，方便快速選擇店鋪；"],
            [FORM_TABLE],
        )
        structure = parse_requirement_document("YL-主数据管理.docx", None, content)
        cases, gaps_gaps, _gapsskips = _build_minimal_cases(
            structure.full_text, "YL-主数据管理.docx", structure=structure)
        companions = [c for c in cases if (c.custom_fields or {}).get("companion_template")]
        assert companions, "伴生场景用例不应消失"
        for case in companions:
            # 通用建议：专用模块名 + 标签双重标识，绝不塞进业务大模块
            assert case.module == GENERIC_SUGGESTION_MODULE
            assert case.custom_fields.get("module_source") == "generic_suggestion"
            assert "通用建议非PRD条款" in case.tags
            assert case.custom_fields.get("coverage_source") == "generic_suggestion"
        # 业务用例不受影响
        assert any(c.module == "店鋪主數據管理" for c in cases)


class TestPendingClassification:
    def test_document_ambiguity_is_tagged_and_stripped_of_assertions(self):
        cases = [
            TestCaseCreate(
                name="店鋪編號-边界值与格式校验",
                module="店鋪主數據管理",
                test_case_steps=[TestStepCreate(step="填寫店鋪編號", result="可提交")],
                custom_fields={
                    "requirement_rule_id": "FR-test",
                    "requirement_rule": {
                        "field": "店鋪編號",
                        "evidence": "未填寫店鋪編號，系統自動生成品牌編號。",
                    },
                    "assertions": [{"kind": "max_length"}],
                },
            )
        ]
        from app.api.v2.ai import _assign_case_metadata

        out = _assign_case_metadata(cases, "未填寫店鋪編號，系統自動生成品牌編號。",
                                    "YL-主数据管理.pdf", None)
        cf = out[0].custom_fields
        assert cf["confirmation_status"] == "pending"
        assert cf["ambiguity_kind"] == AMBIGUITY_DOCUMENT
        assert "assertions" not in cf
        assert out[0].name.startswith("待确认-")

    def test_rule_without_ambiguity_is_not_marked_pending(self):
        """筛选区同名字段规则本身无歧义 → 不得被别的规则的歧义连带标成待确认。"""
        cases = [
            TestCaseCreate(
                name="店鋪編號-边界值与格式校验",
                module="店鋪主數據管理",
                custom_fields={
                    "requirement_rule_id": "FR-filter",
                    "requirement_rule": {
                        "field": "店鋪編號",
                        "evidence": "小於等於30個字符默認為空，為空時查詢全部",
                    },
                },
            )
        ]
        from app.api.v2.ai import _assign_case_metadata

        out = _assign_case_metadata(
            cases,
            "未填寫店鋪編號，系統自動生成品牌編號。\n小於等於30個字符默認為空，為空時查詢全部",
            "YL-主数据管理.pdf", None)
        assert out[0].custom_fields.get("confirmation_status") is None

    def test_manual_supplement_has_its_own_kind(self):
        text = ("1. 導入門店\n導入規則\n"
                "9. 特殊字段：同一个字段支持传多组值时，第一組為必填；每組數據單獨檢驗完整性\n"
                "a. 營業時間1-周幾？\n")
        cases = _build_import_supplement_cases(text, "YL-主数据管理.pdf")
        assert len(cases) == 1
        cf = cases[0].custom_fields
        assert cf["ambiguity_kind"] == AMBIGUITY_MANUAL_SUPPLEMENT
        assert cf["confirmation_status"] == "pending"
        assert cf["ambiguity_reason"], "人工补充项必须写明缺什么，不能用文档歧义的理由"
        assert "歧义" not in cf["ambiguity_reason"] or "未約定" in cf["ambiguity_reason"]

    def test_no_manual_supplement_when_document_defines_overlap(self):
        text = ("導入規則\n"
                "9. 特殊字段：同一个字段支持传多组值时，第一組為必填；營業時間重疊時合併處理\n")
        assert _build_import_supplement_cases(text, "需求.pdf") == []


class TestProductConfirmation:
    """产品确认回填：已裁决的歧义不再重复标待确认，并落成**可判定预期**。"""

    DOC = "小杨生煎-订单评价(问卷调查)功能 (1).pdf"
    #: 文档里同时存在正文 30min 与批注 20min（歧义来源）
    TEXT = (
        "1. 需求背景\n"
        "启用后的问卷将在用户支付完成30min（配置文件）后在订单详情显示“评价”入口"
        "（若订单取消，则不显示）\n"
        "批注：支付完成20min（时间oss配置）\n"
    )

    def _run(self, point: str, name: str) -> TestCaseCreate:
        from app.api.v2.ai import _assign_case_metadata

        case = TestCaseCreate(
            name=name,
            module="营销云-调查问卷管理",
            test_case_steps=[
                TestStepCreate(step=f"准备并执行需求场景：{point}", result=f"核对文档约定：{point}"),
                TestStepCreate(step="检查页面、数据和状态变化，记录实际结果",
                               result=f"实际结果符合原文：{point}"),
            ],
            custom_fields={"requirement_rule_id": "RQ-test", "source_point": point},
        )
        _assign_case_metadata([case], self.TEXT, self.DOC, None)
        return case

    def test_confirmed_ambiguity_is_released_with_determinate_expectation(self):
        point = ("启用后的问卷将在用户支付完成30min（配置文件）后在订单详情显示“评价”入口"
                 "（若订单取消，则不显示）")
        case = self._run(point, point[:120])
        cf = case.custom_fields
        assert cf["confirmation_status"] == "confirmed"
        assert cf["confirmation_key"] == "PR2-EVAL-ENTRY-DELAY"
        # 结论原文照录，不改写、不补充业务规则
        assert "以正文 30min 为准" in cf["confirmed_conclusion"]
        assert cf["confirmed_by"] and cf["confirmed_at"] and cf["confirmation_source"]
        # Q-09a 追加确认的配置落点同样落库
        assert "前端配置文件/常量" in cf["confirmation_config_landing"]
        assert cf["ambiguity_kind"] == AMBIGUITY_DOCUMENT
        # 不再带"待确认-"前缀，并给出可判定预期
        assert not case.name.startswith("待确认-")
        assert "已确认" in case.tags
        assert "待确认" not in case.tags
        assert case.test_case_steps[-1].result == (
            "按正文 30min 核对：前端配置默认 30min（配置文件/常量，无配置界面），"
            "支付完成满 30min 后订单详情出现「评价」入口；订单已取消时不出现该入口"
        )
        assert "产品已确认" in case.description
        # Q-09a 已追加确认 → 不再有未决子问题
        assert not cf.get("confirmation_open_questions")

    def test_same_ambiguity_without_registry_entry_stays_pending(self):
        """同一歧义理由下的**未裁决**用例必须仍待确认（Q-10 订阅消息推送）。"""
        point = "订单支付完成后30分钟（前提是订单非“取消”状态），推送订阅消息"
        case = self._run(point, point)
        cf = case.custom_fields
        assert cf["confirmation_status"] == "pending"
        assert case.name.startswith("待确认-")
        assert "待确认" in case.tags
        assert "已确认" not in case.tags

    def test_registry_requires_document_and_case_match(self):
        from app.services.product_confirmations import find_confirmation

        reason = "正文30分钟与批注20分钟不一致，批注不作为正式规则，时点需确认"
        assert find_confirmation(self.DOC, reason, "订单详情显示“评价”入口") is not None
        assert find_confirmation("别的文档.pdf", reason, "订单详情显示“评价”入口") is None
        assert find_confirmation(self.DOC, reason, "推送订阅消息") is None
        assert find_confirmation(self.DOC, "别的歧义理由", "订单详情显示“评价”入口") is None

    def test_registry_entries_are_complete_and_valid(self):
        import re as _re

        from app.services.product_confirmations import CONFIRMATIONS

        assert CONFIRMATIONS, "登记表不应为空"
        for item in CONFIRMATIONS:
            for field in ("key", "document", "ambiguity_match", "case_match",
                          "conclusion", "confirmed_by", "confirmed_at", "source"):
                assert item.get(field), f"{item.get('key')} 缺字段 {field}"
            _re.compile(item["ambiguity_match"])
            _re.compile(item["case_match"])


class TestParsingDefectFixes:
    """P0 系列解析缺陷的回归（2026-10-14 实测发现，逐条锁定）。"""

    def test_max_count_units_group_dao_day_are_extracted(self):
        from app.services.requirement_structure import extract_constraints

        def kinds(text):
            return {c.kind: c for c in extract_constraints(text)}

        for text, number, unit in (
            ("最多支持選擇7組，", 7.0, "組"),
            ("最多選擇30天", 30.0, "天"),
            ("限制最多设置10道", 10.0, "道"),
            ("最多支持設置2個時間段，且2個時間段不能重複", 2.0, "時間段"),
        ):
            got = kinds(text)
            assert "max_count" in got, f"{text!r} 未识别为数量上限"
            assert got["max_count"].numbers == [number]
            assert got["max_count"].detail.get("unit") == unit

        # 长度类不能被误判成数量（"个字"的"个"不是计量单位）
        length_only = kinds("最多输入140个字")
        assert "max_length" in length_only and "max_count" not in length_only

    def test_option_scope_typo_is_tolerated(self):
        from app.services.requirement_structure import extract_constraints

        # 文档原文错别字「選項發範圍」——不认就会整条筛选字段无用例
        got = {c.kind: c for c in extract_constraints("選項發範圍：待開始、進行中、已完成、失敗")}
        assert got["options"].detail["options"] == ["待開始", "進行中", "已完成", "失敗"]

    def test_condition_phrases_without_value_colon(self):
        from app.services.requirement_structure import extract_condition

        assert extract_condition("當店鋪狀態開啟時，顯示該字段支持設置營業狀態關閉日期") == "店鋪狀態開啟"
        assert extract_condition("排隊打開顯示該字段選項範圍：GULU、QRS") == "排隊打開"
        # rev184：无冒号的繁体「X選擇值」式也被识别，且与「X选择：值」同走
        # 「X=值」归一化——同一语义形态（繁简/有无冒号）产出统一口径。
        assert extract_condition(
            "排隊打開且排隊集成方選擇GULU、QRS時顯示該字段選項範圍：rolling"
        ) == "排隊打開且排隊集成方=GULU、QRS"
        # 「X选择：值」式仍走归一化（既有行为不变）
        assert extract_condition(
            "当变动类型选择：企业储值充值）且变更方式选中：批量导入，显示该字段1、下载模板"
        ) == "变动类型=企业储值充值 且 变更方式=批量导入"
        # 非条件句式不得误判
        assert extract_condition("固定显示，获取用户昵称，为空时显示“-”") == ""

    def test_column_shift_recovers_displaced_field_name_and_drops_group_titles(self):
        content = _build_docx(
            ["1. 店鋪主數據管理", "1.2.1.2 字段說明"],
            [[
                ["字段", "字段类型", "是否必填", "字段说明"],
                # 实测形态：字段名落在"字段类型"列、名字列为空
                ["", "營業時間/選擇組件", "", "每組中需要設置營業時間，一天最多支持設置2個時間段"],
                # 实测形态：分组标题行（只有名称、其余列全空）
                ["排隊/控制該店鋪是否支持排隊，以及優先取票規則", "", "", ""],
                ["排隊", "開關", "是", "默認為關閉"],
            ]],
        )
        structure = parse_requirement_document("YL.docx", None, content)
        names = [r.name for r in structure.field_rules]
        assert "排隊" in names
        assert any("營業時間" in n for n in names), names
        assert not any("控制該店鋪是否支持" in n for n in names), names
        assert any("分组标题" in item["reason"] for item in structure.rejected_fields)

    def test_export_detail_table_gets_integrity_case(self):
        content = _build_docx(
            ["2.3.6 導出明細", "導出明細格式："],
            [[
                ["字段", "说明"],
                ["填写渠道", "固定显示，包括支付宝小程序/微信小程序"],
                ["用户名", "固定显示，获取用户昵称，为空时显示“-”"],
            ]],
        )
        structure = parse_requirement_document("PR2.docx", None, content)
        cases, gaps_gaps, _gapsskips = _build_minimal_cases(structure.full_text, "PR2.docx", structure=structure)
        export = [c for c in cases if (c.custom_fields or {}).get("export_detail_case")]
        assert export, [c.name for c in cases]
        assert "導出明細" in export[0].name and "欄位完整性" in export[0].name
        steps = " ".join(s.step + " " + s.result for s in export[0].test_case_steps)
        assert "填写渠道" in steps and "用户名" in steps
        assert "-" in steps  # 空值显示约定被写进步骤

    def test_source_location_and_fingerprint_are_recorded(self):
        content = _build_docx(
            ["1. 店鋪主數據管理", "1.2.1.2 字段說明", "維護店鋪的主數據，方便快速選擇店鋪；"],
            [[["字段", "字段类型", "是否必填", "字段说明"],
              ["店鋪編號", "輸入框", "否", "小於等於30個字符；字符類型：大小寫字母、數字"]]],
        )
        structure = parse_requirement_document("YL.docx", None, content)
        cases, gaps_gaps, _gapsskips = _build_minimal_cases(structure.full_text, "YL.docx", structure=structure)
        assert cases
        # 指纹：同一文本稳定、不同文本不同
        digests = {(c.custom_fields or {}).get("source_text_sha1") for c in cases}
        assert len(digests) == 1 and None not in digests
        other, _gaps, _skips = _build_minimal_cases(
            "B2B储值订单不能退款，储值订单不在C端显示；", "YL.docx", structure=None)
        assert other and (other[0].custom_fields or {}).get("source_text_sha1") not in digests
        # 精确来源：字段规则用例带表行，其余用例也应尽量带行号
        assert all(c.source_location for c in cases if "字段规则" in (c.tags or []))


class TestTitleDeduplication:
    @staticmethod
    def _case(rule_id: str, table: int, row: int) -> TestCaseCreate:
        return TestCaseCreate(
            name="店鋪編號-边界值与格式校验",
            module="店鋪主數據管理",
            test_case_steps=[TestStepCreate(step="填寫店鋪編號", result="可提交")],
            custom_fields={
                "requirement_rule_id": rule_id,
                "requirement_rule": {"field": "店鋪編號", "source": {"table": table, "row": row}},
            },
        )

    def test_identical_rule_is_deduplicated(self):
        a, b = self._case("FR-1", 1, 6), self._case("FR-1", 1, 6)
        kept = _resolve_duplicate_titles([a, b])
        assert len(kept) == 1
        assert kept[0].custom_fields["duplicate_removed"] == 1

    def test_different_rules_are_renamed_not_dropped(self):
        a, b = self._case("FR-1", 1, 6), self._case("FR-2", 8, 2)
        kept = _resolve_duplicate_titles([a, b])
        assert len(kept) == 2
        assert len({c.name for c in kept}) == 2
        assert {c.custom_fields["requirement_rule_id"] for c in kept} == {"FR-1", "FR-2"}
        for case in kept:
            assert case.custom_fields.get("title_disambiguation")

    def test_label_falls_back_to_module_then_rule_id(self):
        """场景标签仍重名时升级为「模块-场景」，再重名则追加规则 ID（保证标题唯一）。"""
        a = self._case("FR-1", 1, 6)
        b = self._case("FR-2", 1, 6)  # 同模块同表同行 → 标签也相同
        kept = _resolve_duplicate_titles([a, b])
        assert len(kept) == 2
        names = [c.name for c in kept]
        assert len(set(names)) == 2, names
        assert any("FR-1" in n or "FR-2" in n for n in names), names
        assert all(c.custom_fields.get("title_disambiguation") for c in kept)

    def test_degenerate_case_without_any_evidence_still_unique(self):
        """连表号/行号都没有时，标题仍需唯一（退回"来源位置待补"→ 追加规则 ID）。"""
        a = TestCaseCreate(name="X-校验", module="M", custom_fields={"requirement_rule_id": "FR-a"})
        b = TestCaseCreate(name="X-校验", module="M", custom_fields={"requirement_rule_id": "FR-b"})
        kept = _resolve_duplicate_titles([a, b])
        assert len(kept) == 2 and len({c.name for c in kept}) == 2
