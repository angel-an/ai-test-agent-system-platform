# -*- coding: utf-8 -*-
"""3a 回归：PDF 字段说明表抽取不得静默归零。

两个真实缺陷（2026-09-23/24 现场，PR-9 `CS Module方案説明2.0….pdf`）：

1. **表结构没被识别**：PyMuPDF 对无边框线的表格 `page.find_tables()` 返回 0 张表
   → `field_rules=0` → 整批用例 0 条可执行断言（`test_borderless_pdf_*`）；
2. **列错位/整行粘连**：抽取把一行粘成一格，说明列文字落到"字段名"列被拒绝后整行丢弃
   → `field_rules=0`（`test_glued_row_*`，样本取自 `test_generation_drift_gates.py`
   记录的真实形态）。

修复目标是"不静默归零"，**不是**把段落文字伪装成字段：`test_paragraph_*` 守住这条底线。
"""

from __future__ import annotations

import fitz
import pytest

from app.services.requirement_structure import (
    _split_glued_row,
    _split_tables,
    _tables_from_text_lines,
    parse_requirement_document,
)

FIELD_HEADER = ["字段名稱", "類型", "必填", "説明"]

#: PR-9 实测形态：整行被抽成一格，多列用**单空格**粘在一起
GLUED_ROW = ["Member 是 必填，選取會員名稱Case  是 1. 如在GoodWill模塊創建的，必選選擇case"]
#: 纯段落行（不得被伪装成字段）
PARAGRAPH_ROW = ["3. 點擊Case Subcategory Level1名稱，跳轉到詳情"]


def _borderless_pdf(lines: list[str]) -> bytes:
    """生成一份"有表格语义但无边框线"的 PDF（find_tables() 认不出来）。"""
    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for line in lines:
        page.insert_text((50, y), line, fontname="china-s", fontsize=10)
        y += 16
    payload = doc.tobytes()
    doc.close()
    return payload


class TestBorderlessPdfFieldTable:
    LINES = [
        "2.1 字段说明",
        "字段名稱      類型      必填      説明",
        "Case Title    Text      是        導出標題格式：cs_cases_YYYYMMDD_hhmmss",
        "Case Status   Select    是        選項範圍：New, Rejected, Closed",
        "Modified On   Date      否        最新修改時間",
        "",
        "2.2 其他说明",
        "本節描述匯出邏輯。",
    ]

    def test_find_tables_returns_nothing_for_borderless_table(self):
        """先钉住前提：这类 PDF 的表格结构确实识别不出来（否则本测试无意义）。"""
        doc = fitz.open(stream=_borderless_pdf(self.LINES), filetype="pdf")
        page = doc[0]
        assert len(getattr(page.find_tables(), "tables", []) or []) == 0
        doc.close()

    def test_field_rules_are_recovered_from_text_lines(self):
        structured = parse_requirement_document(
            "CS Module方案説明2.0.pdf", "application/pdf", _borderless_pdf(self.LINES)
        )
        summary = structured.to_dict()
        assert summary["source_type"] == "pdf"
        assert summary["rebuilt_table_count"] >= 1
        assert summary["field_rule_count"] == 3, summary
        names = [rule.name for rule in structured.field_rules]
        assert names == ["Case Title", "Case Status", "Modified On"]
        # 列映射正确（不是把说明文字当字段名）
        by_name = {rule.name: rule for rule in structured.field_rules}
        assert by_name["Case Title"].field_type == "Text"
        assert by_name["Case Title"].required is True
        assert by_name["Modified On"].required is False

    def test_no_field_table_means_no_rules(self):
        """没有字段说明表时不得凭空产出规则。"""
        structured = parse_requirement_document(
            "plain.pdf",
            "application/pdf",
            _borderless_pdf(["1. 背景", "本需求描述匯出邏輯，沒有字段說明表。"]),
        )
        assert structured.field_rules == []
        assert structured.rebuilt_table_count == 0


class TestGluedRowRecovery:
    def test_glued_row_recovers_field_label(self):
        rules, _templates, rejected, rebalanced, _disp = _split_tables([[FIELD_HEADER, GLUED_ROW]])
        assert [rule.name for rule in rules] == ["Member"]
        assert rebalanced == 1
        assert rejected == []
        # 恢复出的规则留有审计痕迹
        assert rules[0].source["rebalanced"] is True
        # 说明列保留原文（供 extract_constraints 抽断言）
        assert rules[0].description.startswith("是 必填")

    def test_glued_row_does_not_invent_type(self):
        """类型列一律留空：不猜类型。

        required 不再断言为 False——该行的原文里**确实写着**"必填"
        （`Member 是 必填，…`），据此判定 required=True 是文档依据，不是编造。
        """
        rules, _t, _r, _rb, _disp = _split_tables([[FIELD_HEADER, GLUED_ROW]])
        assert rules[0].field_type == ""
        assert rules[0].required is True

    def test_two_space_glued_row_is_split_positionally(self):
        row = ["Field A  Text  是  desc A"]
        rules, _t, rejected, rebalanced, _disp = _split_tables([[FIELD_HEADER, row]])
        assert [rule.name for rule in rules] == ["Field A"]
        assert rebalanced == 1
        assert rejected == []

    def test_paragraph_row_is_still_rejected(self):
        """底线：段落文字不得被伪装成字段（不静默丢弃，但也不产出规则）。"""
        rules, _t, rejected, _rb, _disp = _split_tables([[FIELD_HEADER, PARAGRAPH_ROW]])
        assert rules == []
        assert len(rejected) == 1
        assert rejected[0]["name"].startswith("3. 點擊Case")

    def test_split_glued_row_returns_none_without_marker(self):
        assert _split_glued_row("只是普通的一段說明文字", 4) is None
        assert _split_glued_row("Member 是 必填", 1) is None


class TestCompilabilityWarning:
    """字段规则为 0 时必须显式告警（PR-9 实测：36 条用例 0 断言却无任何提示）。"""

    def test_warns_when_no_field_rules(self):
        from app.api.v2.ai import _build_compilability_warning

        warning = _build_compilability_warning(
            {"source_type": "pdf", "field_rule_count": 0, "rebuilt_table_count": 0,
             "rebalanced_rows": 7, "rejected_field_count": 7},
            36,
        )
        assert warning is not None
        assert "field_rule_count=0" in warning
        assert "36 条用例不含可执行断言" in warning
        assert "重建表格数=0" in warning and "重列化行数=7" in warning

    def test_silent_when_field_rules_exist(self):
        from app.api.v2.ai import _build_compilability_warning

        assert _build_compilability_warning({"field_rule_count": 17}, 53) is None

    def test_empty_structure_still_warns(self):
        from app.api.v2.ai import _build_compilability_warning

        assert _build_compilability_warning({}, 0) is not None

    def test_degraded_parse_is_distinguished_from_missing_table(self):
        """解析失败降级 ≠ 文档没有字段表：两者的排查方向完全不同。"""
        from app.api.v2.ai import _build_compilability_warning

        warning = _build_compilability_warning(
            {"source_type": "text", "field_rule_count": 0,
             "degraded": True, "degraded_reason": "ValueError: bad pdf"},
            36,
        )
        assert warning is not None
        assert "结构化解析失败" in warning
        assert "bad pdf" in warning
        assert "不是文档缺字段说明表" in warning


class TestExtractStructureDegradation:
    def test_failure_marks_degraded_and_keeps_text(self, monkeypatch):
        """解析抛异常时：不得静默——必须留痕并保留纯文本。"""
        from app.api.v2 import ai as ai_module

        def _boom(*_args, **_kwargs):
            raise ValueError("bad pdf")

        monkeypatch.setattr(ai_module, "parse_requirement_document", _boom)
        structured = ai_module._extract_structure(
            "broken.pdf", "application/pdf", b"%PDF-1.4 fake"
        )
        summary = structured.to_dict()
        assert summary["degraded"] is True
        assert "bad pdf" in summary["degraded_reason"]
        assert summary["field_rule_count"] == 0
        assert structured.narrative_text  # 纯文本兜底仍在

    def test_success_is_not_marked_degraded(self):
        from app.api.v2 import ai as ai_module

        structured = ai_module._extract_structure(
            "ok.txt", "text/plain", "1. 需求：本需求描述匯出邏輯。".encode("utf-8")
        )
        assert structured.to_dict()["degraded"] is False


class TestConstraintVocabularyCsSpec:
    """真件（CS Module 方案説明 2.0.pdf）的约束措辞必须能被抽取。

    这是 PR-9 "0 条规则用例" 的**真正根因**：128 条字段规则里 `constraints` 全为空、
    `required` 全为 False，于是生成器（正确地）拒绝编造断言 → 0 条规则用例。
    不是表格识别问题，是**词表不匹配繁體/列表式表述**。
    """

    def test_required_written_in_description(self):
        from app.services.requirement_structure import required_from_text

        assert required_from_text("1. 必填2. 限制輸入50個字符3. 未輸入提示：請輸入工單標題") is True
        assert required_from_text("非必選内容，上傳文件不限格式，上傳文件大小不超過10m") is False
        assert required_from_text("非必選内容，但附件必填") is True
        assert required_from_text("工單編號") is False

    def test_max_length_traditional_list_style(self):
        from app.services.requirement_structure import extract_constraints

        found = {c.kind: c for c in extract_constraints("1. 必填2. 限制輸入50個字符3. 未輸入提示：請輸入工單標題")}
        assert "max_length" in found
        assert found["max_length"].numbers == [50.0]

    def test_options_from_enumeration_wording(self):
        from app.services.requirement_structure import extract_constraints

        found = {c.kind: c for c in extract_constraints(
            "服務模式展示枚舉内容Dine-in，In-Store Takeaway，Onlinetakeaway，3rd Party"
        )}
        assert "options" in found
        assert found["options"].detail["options"] == [
            "Dine-in", "In-Store Takeaway", "Onlinetakeaway", "3rd Party",
        ]

    def test_options_from_enumeration_colon_wording(self):
        """CS 規格书 Status 写作「枚舉：Active，Closed，Overdue」——冒号式枚举也要能抽取。

        背景（2026-09-28 量化）：128 条字段规则里 109 条无约束，其中仅 3 条原文含约束措辞，
        本条是其中**唯一**可直接转成可执行断言的（另外两条是条件可编辑说明与展示格式说明）。
        原词表只有「枚舉内容/枚舉值」形态 → 该枚举完全没被抽取，Status 字段生成不出断言。
        """
        from app.services.requirement_structure import extract_constraints

        found = {c.kind: c for c in extract_constraints("枚舉：Active，Closed，Overdue")}
        assert found["options"].detail["options"] == ["Active", "Closed", "Overdue"]

        simp = {c.kind: c for c in extract_constraints("枚举: Active、Closed")}
        assert simp["options"].detail["options"] == ["Active", "Closed"]

        # 裸「枚舉」不带冒号仍不进词表（正文普通提及不得当成枚举声明）
        assert "options" not in {c.kind for c in extract_constraints("枚舉字段說明如下")}

    def test_default_value_wording(self):
        from app.services.requirement_structure import extract_constraints

        found = {c.kind: c for c in extract_constraints("創建工單后，默認為In Progress")}
        assert found["default"].detail["value"] == "In Progress"

        mixed = {c.kind: c for c in extract_constraints("默認YL選項，支持選擇ETZ")}
        assert mixed["default"].detail["value"] == "YL"
        assert mixed["options"].detail["options"] == ["ETZ"]

    def test_file_size_with_negation_word(self):
        from app.services.requirement_structure import extract_constraints

        found = {c.kind: c for c in extract_constraints(
            "非必選内容，上傳文件不限格式，上傳文件大小不超過10m，上傳后顯示文件標題"
        )}
        assert found["file_size"].numbers == [10.0]

    def test_legacy_simplified_phrasings_still_work(self):
        """PR-2 文档的简体表述不得回归（否则是拿新客户换旧客户）。"""
        from app.services.requirement_structure import extract_constraints

        assert {c.kind for c in extract_constraints("最多输入10个手机号")} == {"max_count"}
        assert {c.kind for c in extract_constraints("选项范围：企业储值充值、储值退款")} == {"options"}
        assert {c.kind for c in extract_constraints("默认选中：企业储值充值")} == {"default"}
        assert {c.kind for c in extract_constraints("文件格式仅支持：xlsx、csv")} == {"formats"}
        assert {c.kind for c in extract_constraints("小数点后最多保留2位")} == {"decimal_places"}
        assert {c.kind for c in extract_constraints("范围：(0,99999]")} == {"range"}

    def test_prose_tab_word_is_not_an_option(self):
        """裸「選項」不进词表：正文里的"選項卡"不得被当成枚举值。"""
        from app.services.requirement_structure import extract_constraints

        assert "options" not in {c.kind for c in extract_constraints("點擊選項卡切換到下一頁")}

    def test_field_rule_gets_required_from_description(self):
        """端到端：说明里写必填 → FieldRule.required 为真（表格列没有必填列也能取到）。"""
        header = ["字段名稱", "説明"]
        row = ["Case Title", "1. 必填2. 限制輸入50個字符3. 未輸入提示：請輸入工單標題"]
        rules, _t, _r, _rb, _disp = _split_tables([[header, row]])
        assert [rule.name for rule in rules] == ["Case Title"]
        assert rules[0].required is True
        assert {c.kind for c in rules[0].constraints} == {"max_length"}


class TestTextLineReconstruction:
    def test_reconstructs_table_from_lines(self):
        text = "\n".join([
            "字段名稱      類型      必填      説明",
            "Case Title    Text      是        標題格式",
            "Case Status   Select    否        狀態",
        ])
        tables = _tables_from_text_lines(text)
        assert len(tables) == 1
        assert tables[0][0] == ["字段名稱", "類型", "必填", "説明"]
        assert len(tables[0]) == 3

    def test_ignores_prose_without_field_header(self):
        assert _tables_from_text_lines("這是一段沒有表頭的需求描述文字。\n另一段。") == []

    def test_stops_at_single_segment_line(self):
        text = "\n".join([
            "字段名稱    類型    必填    説明",
            "Case Title  Text    是      標題格式",
            "3. 下一節標題",
            "正文段落內容。",
        ])
        tables = _tables_from_text_lines(text)
        assert len(tables) == 1
        assert len(tables[0]) == 2  # 表头 + 1 行数据，不吞正文

    @pytest.mark.parametrize("suffix", [".pdf"])
    def test_docx_path_unchanged(self, suffix):
        """docx 仍走真实表格对象（PR-2 现场 17 条字段规则），回归不受影响。"""
        from docx import Document

        doc = Document()
        doc.add_paragraph("2.1 字段说明")
        table = doc.add_table(rows=3, cols=4)
        for index, value in enumerate(["字段名稱", "類型", "必填", "説明"]):
            table.cell(0, index).text = value
        for index, value in enumerate(["企业账户充值金额", "Input", "是", "范围：(0,99999]"]):
            table.cell(1, index).text = value
        for index, value in enumerate(["备注", "Textarea", "否", "最多输入 200 个字"]):
            table.cell(2, index).text = value
        import io

        buffer = io.BytesIO()
        doc.save(buffer)
        structured = parse_requirement_document(
            f"需求{suffix}", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            buffer.getvalue(),
        )
        names = [rule.name for rule in structured.field_rules]
        assert names == ["企业账户充值金额", "备注"]
        assert structured.field_rules[0].required is True


def test_narrow_header_continuation_table_no_crash():
    """窄表头（宽度<3）续表不得 IndexError（rev185 干跑：CS 方案书真实 PDF 崩溃）。"""
    from app.services.requirement_structure import _continuation_offset

    header = ["字段", "字段说明"]
    rows = [
        ["状态", "启用后显示"],
        ["备注", "选填"],
    ]
    # 修复前：cells[off+2] 越行尾 → IndexError
    assert _continuation_offset(rows, header) is None
